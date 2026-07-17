"""Tests multi-sources (V1.0) : MySQL et CSV/Excel via la couche d'abstraction.

Les tests MySQL sont ignorés si aucune instance MySQL/MariaDB n'est joignable.
Les tests fichiers (CSV/Excel) tournent hors-ligne.
"""
from __future__ import annotations

import csv
import os

import pytest

from app.services.sources.base import SourceConfig
from app.services.sources.files import FileAdapter

MYSQL = {
    "engine": "mysql", "host": "127.0.0.1", "port": 3306,
    "database": "noreon_demo_mysql", "username": "noreon_ro", "password": "readonly",
}


def _mysql_available() -> bool:
    try:
        import pymysql
        c = pymysql.connect(host=MYSQL["host"], port=MYSQL["port"], user=MYSQL["username"],
                            password=MYSQL["password"], database=MYSQL["database"], connect_timeout=3)
        c.close()
        return True
    except Exception:
        return False


mysql_required = pytest.mark.skipif(not _mysql_available(), reason="MySQL/MariaDB indisponible.")


# ---------------------------------------------------------------------------
# Petits objets « table/colonne » pour appeler profile_table hors ORM.
# ---------------------------------------------------------------------------
class _T:
    def __init__(self, schema, name, rows):
        self.schema_name, self.table_name, self.estimated_rows = schema, name, rows


class _C:
    def __init__(self, name, dtype):
        self.name, self.data_type = name, dtype


def _cols(table_info):
    return [_C(c.name, c.data_type) for c in table_info.columns]


# ---------------------------------------------------------------------------
# CSV / Excel (hors-ligne)
# ---------------------------------------------------------------------------
@pytest.fixture
def csv_adapter(tmp_path, monkeypatch):
    monkeypatch.setenv("NOREON_DATA_DIR", str(tmp_path))
    path = tmp_path / "ventes.csv"
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "produit", "categorie", "montant", "email"])
        for i in range(1, 61):
            w.writerow([i, f"P{i}", ["Textile", "Tech", "Maison"][i % 3],
                        round(10 + i * 2.5, 2), f"c{i}@ex.com" if i % 8 else "invalide"])
    return FileAdapter(SourceConfig(engine="csv", file_path=str(path), options={"format": "csv"}))


def test_csv_read_only_and_scan(csv_adapter):
    assert csv_adapter.check_read_only()["read_only"] is True
    assert csv_adapter.test_connection()["ok"] is True
    scan = csv_adapter.introspect()
    table = next(t for t in scan.tables if t.name == "ventes")
    assert table.estimated_rows == 60
    assert {c.name for c in table.columns} == {"id", "produit", "categorie", "montant", "email"}
    # Inférence de type : montant en REAL, id en INTEGER.
    types = {c.name: c.data_type for c in table.columns}
    assert types["id"] == "integer" and types["montant"] == "real"


def test_csv_query_with_guardrails(csv_adapter):
    r = csv_adapter.run_query(
        "SELECT categorie, count(*) n, avg(montant) m FROM ventes GROUP BY categorie", connection_id=1
    )
    assert set(r.columns) == {"categorie", "n", "m"}
    assert sum(row[1] for row in r.rows) == 60
    assert "LIMIT" in r.guarded_sql.upper()


def test_csv_write_blocked(csv_adapter):
    from app.services.sql_guard import SQLGuardError

    with pytest.raises(SQLGuardError):
        csv_adapter.run_query("DELETE FROM ventes", connection_id=1)


def test_csv_profiling_detects_pii_and_invalid(csv_adapter):
    from app.services.profiler import profile_table

    scan = csv_adapter.introspect()
    ventes = next(t for t in scan.tables if t.name == "ventes")
    t = _T("main", "ventes", 60)
    profiles = profile_table(csv_adapter, t, _cols(ventes))
    by = {p.column_name: p for p in profiles}
    assert by["email"].pii_type == "email"
    # Quelques emails « invalide » → validité < 100%.
    assert by["email"].invalid_count and by["email"].invalid_count > 0
    assert by["id"].distinct_count == 60


def test_excel_adapter(tmp_path, monkeypatch):
    monkeypatch.setenv("NOREON_DATA_DIR", str(tmp_path))
    from openpyxl import Workbook

    path = tmp_path / "classeur.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "clients"
    ws.append(["id", "nom", "ville"])
    for i in range(1, 21):
        ws.append([i, f"Client {i}", ["Paris", "Lyon"][i % 2]])
    wb.save(path)

    adapter = FileAdapter(SourceConfig(engine="excel", file_path=str(path), options={"format": "excel"}))
    assert adapter.test_connection()["ok"] is True
    scan = adapter.introspect()
    clients = next(t for t in scan.tables if t.name == "clients")
    assert clients.estimated_rows == 20
    r = adapter.run_query("SELECT count(*) FROM clients", connection_id=2)
    assert r.rows[0][0] == 20


# ---------------------------------------------------------------------------
# MySQL / MariaDB (base réelle)
# ---------------------------------------------------------------------------
@pytest.fixture
def mysql_adapter():
    from app.services.sources.mysql import MySQLAdapter

    return MySQLAdapter(SourceConfig(**MYSQL))


@mysql_required
def test_mysql_read_only(mysql_adapter):
    ro = mysql_adapter.check_read_only()
    assert ro["read_only"] is True


@mysql_required
def test_mysql_scan_declared_and_inferred_fk(mysql_adapter):
    scan = mysql_adapter.introspect()
    names = {t.name for t in scan.tables}
    assert {"customers", "orders", "stores"} <= names
    kinds = {(r.from_table, r.from_column, r.kind) for r in scan.relations}
    assert ("orders", "customer_id", "declared") in kinds
    assert ("customers", "store_id", "inferred") in kinds


@mysql_required
def test_mysql_query_guardrails(mysql_adapter):
    r = mysql_adapter.run_query("SELECT count(*) n FROM customers", connection_id=3)
    assert r.rows[0][0] == 100
    assert "LIMIT" in r.guarded_sql.upper()


@mysql_required
def test_mysql_write_blocked(mysql_adapter):
    from app.services.sql_guard import SQLGuardError

    with pytest.raises(SQLGuardError):
        mysql_adapter.run_query("UPDATE customers SET city='x'", connection_id=3)


@mysql_required
def test_mysql_profiling(mysql_adapter):
    from app.services.profiler import profile_table

    scan = mysql_adapter.introspect()
    cust = next(t for t in scan.tables if t.name == "customers")
    t = _T(MYSQL["database"], "customers", cust.estimated_rows)
    profiles = profile_table(mysql_adapter, t, _cols(cust))
    by = {p.column_name: p for p in profiles}
    assert by["email"].pii_type == "email"
    assert by["email"].null_rate and by["email"].null_rate > 0
